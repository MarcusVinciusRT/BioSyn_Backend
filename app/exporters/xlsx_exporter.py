"""Relatorio em XLSX."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.schemas.comum import para_fuso_local
from app.schemas.relatorio import RelatorioResponse

CABECALHO = ["Métrica", "Filtro", "Valor", "Unidade"]
LARGURAS = (38, 22, 16, 14)


def gerar_xlsx(relatorio: RelatorioResponse) -> bytes:
    planilha = Workbook()
    aba = planilha.active
    aba.title = "Relatório"

    periodo = relatorio.periodo
    aba.append([f"Relatório BioSyn — {periodo.inicio:%d/%m/%Y} a {periodo.fim:%d/%m/%Y}"])
    aba["A1"].font = Font(bold=True, size=13)
    # para_fuso_local: sem isso o cabecalho sairia em UTC enquanto o JSON e o
    # nome do arquivo saem no fuso local -- tres horas de diferenca na mesma resposta.
    aba.append([f"Gerado em {para_fuso_local(relatorio.gerado_em):%d/%m/%Y %H:%M}"])
    aba.append([])

    linha_cabecalho = aba.max_row + 1
    aba.append(CABECALHO)
    for coluna in range(1, len(CABECALHO) + 1):
        celula = aba.cell(row=linha_cabecalho, column=coluna)
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="center")

    for resultado in relatorio.resultados:
        aba.append(
            [
                resultado.nome,
                resultado.filtro or "—",
                resultado.valor,
                resultado.unidade,
            ]
        )
        if resultado.unidade == "percentual" and resultado.valor is not None:
            aba.cell(row=aba.max_row, column=3).number_format = '0.0"%"'

    for indice, largura in enumerate(LARGURAS, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura

    # Congela o cabecalho para a tabela continuar legivel ao rolar.
    aba.freeze_panes = aba.cell(row=linha_cabecalho + 1, column=1)

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()
